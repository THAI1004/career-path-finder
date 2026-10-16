from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "TAI_LIEU_MO_TA_HE_THONG_HUONG_NGHIEP.md"
OUT = ROOT / "output" / "word" / "Mo_ta_hoan_chinh_du_an_EduCareer_2026.docx"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(cell, color='D9E2F3'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    b = tcPr.first_child_found_in('w:tcBorders')
    if b is None:
        b = OxmlElement('w:tcBorders')
        tcPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        el = b.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            b.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)

def set_cell_text(cell, text, bold=False, size=8.2, color='1F2937'):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run(re.sub(r'\*\*|`', '', text).strip())
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = 'Aptos'
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, rows):
    if len(rows) < 2:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            set_cell_text(cell, value, bold=i == 0, size=7.4 if len(rows[0]) >= 7 else 8.5, color='FFFFFF' if i == 0 else '1F2937')
            shade(cell, '1F4E78' if i == 0 else ('F3F7FB' if i % 2 == 0 else 'FFFFFF'))
            borders(cell)
    for row in table.rows:
        for cell in row.cells:
            cell.margin_top = 70
            cell.margin_bottom = 70
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_inline_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    # basic markdown emphasis
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('*') and part.endswith('*'):
            r = p.add_run(part[1:-1]); r.italic = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
        else:
            p.add_run(part)
    return p

def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith('|'):
        line = lines[i].strip()
        cells = [c.strip() for c in line.strip('|').split('|')]
        if not all(re.fullmatch(r':?-{3,}:?', c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i

def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)
    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10)
    styles['Normal'].paragraph_format.line_spacing = 1.08
    for name, size, color in [('Title', 20, '1F2937'), ('Heading 1', 15, '1F4E78'), ('Heading 2', 12, '2563EB'), ('Heading 3', 10.5, '374151')]:
        st = styles[name]; st.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
    title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run('Mô tả hoàn chỉnh dự án EduCareer')
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = sub.add_run('Đặc tả chức năng và luồng vận hành theo định hướng tuyển sinh 2026'); rr.italic = True; rr.font.size = Pt(11); rr.font.color.rgb = RGBColor.from_string('4B5563')
    doc.add_paragraph()

    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    i = 0; in_code = False; code = []
    while i < len(lines):
        raw = lines[i]; s = raw.strip()
        if s.startswith('```'):
            if not in_code:
                in_code = True; code = []
            else:
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.2); p.paragraph_format.space_after = Pt(7)
                r = p.add_run('\n'.join(code)); r.font.name = 'Consolas'; r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string('374151')
                in_code = False
            i += 1; continue
        if in_code:
            code.append(raw); i += 1; continue
        if not s or s == '---':
            i += 1; continue
        if s.startswith('|'):
            rows, i = parse_table(lines, i); add_table(doc, rows); continue
        if s.startswith('# '):
            i += 1; continue
        if s.startswith('## '):
            doc.add_paragraph(s[3:].strip(), style='Heading 1'); i += 1; continue
        if s.startswith('### '):
            doc.add_paragraph(s[4:].strip(), style='Heading 2'); i += 1; continue
        if s.startswith('#### '):
            doc.add_paragraph(s[5:].strip(), style='Heading 3'); i += 1; continue
        if s.startswith('> '):
            p = add_inline_paragraph(doc, s[2:]); p.paragraph_format.left_indent = Inches(0.2); p.paragraph_format.right_indent = Inches(0.2)
            for r in p.runs: r.font.color.rgb = RGBColor.from_string('4B5563'); r.italic = True
            i += 1; continue
        if re.match(r'^\d+\.\s+', s):
            text = re.sub(r'^\d+\.\s+', '', s); add_inline_paragraph(doc, text, style='List Number'); i += 1; continue
        if s.startswith('- '):
            add_inline_paragraph(doc, s[2:], style='List Bullet'); i += 1; continue
        add_inline_paragraph(doc, s); i += 1

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = 'Mô tả hoàn chỉnh dự án EduCareer 2026'
    doc.core_properties.subject = 'Đặc tả chức năng và luồng vận hành hệ thống hướng nghiệp'
    doc.core_properties.author = 'EduCareer'
    doc.save(OUT)
    print(OUT)

if __name__ == '__main__':
    build()
